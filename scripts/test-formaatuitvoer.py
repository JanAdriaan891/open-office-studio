"""Onafhankelijke inhoudscontrole van de native conversieproef; vereist python-docx, openpyxl en python-pptx."""
import pathlib
import subprocess
import sys

from docx import Document
from openpyxl import load_workbook
from pptx import Presentation

map = pathlib.Path(sys.argv[1])
origineel = Document("docs/voorbeelden/Ontwerpnotitie.docx")
for route in ("odt", "doc"):
    document = Document(map / f"Ontwerpnotitie-via-{route}.docx")
    tekst = "\n".join(p.text for p in document.paragraphs)
    for alinea in origineel.paragraphs:
        assert not alinea.text or alinea.text in tekst, (route, alinea.text)
    assert len(document.tables) == len(origineel.tables)
    print(f"Writer via {route}: alineatekst en aantal tabellen behouden.")

bronboek = load_workbook("docs/voorbeelden/Materiaaloverzicht.xlsx")
for route in ("ods", "xls"):
    bestand = map / f"Materiaaloverzicht-via-{route}.xlsx"
    boek = load_workbook(bestand)
    waarden = load_workbook(bestand, data_only=True)
    assert boek.active["G17"].value == "=SUM(G7:G14)"
    assert waarden.active["G17"].value == 3068
    for rij in bronboek.active:
        for cel in rij:
            if cel.value is not None:
                assert boek.active[cel.coordinate].value == cel.value, (route, cel.coordinate)
    print(f"Calc via {route}: alle ingevulde cellen, formules en berekend totaal behouden.")

def teksten(presentatie):
    return [p.text for dia in presentatie.slides for vorm in dia.shapes if vorm.has_text_frame for p in vorm.text_frame.paragraphs if p.text]

origineel = Presentation("docs/voorbeelden/Ontwerpvisie.pptx")
for route in ("odp", "ppt"):
    presentatie = Presentation(map / f"Ontwerpvisie-via-{route}.pptx")
    assert len(presentatie.slides) == len(origineel.slides)
    assert sorted(teksten(presentatie)) == sorted(teksten(origineel))
    print(f"Presentatie via {route}: dia-aantal en diatekst behouden.")

for naam, zoektekst in (("Ontwerpnotitie", "Ruimte voor ideeën"), ("Materiaaloverzicht", "Materiaaloverzicht"), ("Ontwerpvisie", "Samen ruimte")):
    tekst = subprocess.check_output(["pdftotext", str(map / f"{naam}.pdf"), "-"], text=True)
    assert zoektekst.casefold() in tekst.casefold(), naam
    print(f"PDF {naam}: leesbare tekst aanwezig.")
